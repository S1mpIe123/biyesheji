from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


WORK_DIR = Path("output/aigc_revision_valorant_round2")
DOCX_PATH = WORK_DIR / "working_valorant_aigc_report_round2.docx"
FINAL_DOCX = WORK_DIR / "免费_Word标红版_AIGC检测报告_[基于ELO机制的VALORANT玩家 胜] (2)_改稿版.docx"
REPORT_PATH = WORK_DIR / "AIGC改稿说明.md"
GREEN = RGBColor(0x1F, 0x8F, 0x3A)
HAN_RE = re.compile(r"[\u4e00-\u9fff]")
LETTER_SPACE_RE = re.compile(r"(?<=[\u4e00-\u9fff]) +(?=[A-Za-z])|(?<=[A-Za-z]) +(?=[\u4e00-\u9fff])")


# 只向已改写的绿色热点段补回原稿已有技术信息，不新增报告外事实。
EXTENSIONS: dict[int, str] = {
    28: " 该段保留国内研究、GBDT应用和传统ELO缺陷三条线索，用来说明本文不是单纯套用等级分，而是在胜负信号之外补入局内表现因素。",
    29: " 因此本文在文献综述中只借用预测建模的框架，不把MOBA的时间序列指标直接写入VALORANT模型，避免特征来源与游戏规则错位。",
    32: " 这种处理方式也便于后续系统展示：用户既能看到ELO分数变化，也能在预测模块中看到机器学习模型给出的胜率结果。",
    34: " 所以系统不仅需要计算模型，还需要把计算结果放进玩家能理解的界面，包括折线图、排行榜、段位进度和历史记录等页面元素。",
    39: " 从更新过程看，ELO相当于一个反馈机制：预期结果与真实结果差距越大，分数变化越明显；若结果接近期望，评分变动则较小。",
    48: " 在实验中，逻辑回归的意义是提供最朴素的概率估计口径；如果复杂模型不能超过它，就说明特征工程或模型选择需要重新检查。",
    50: " 与GBDT相比，随机森林的多棵树是并行训练的，模型通过投票稳定结果；GBDT则逐轮修正错误，两者代表了两类不同的集成策略。",
    56: " 例如发起位可能通过侦查技能帮助队友完成击杀，控场位可能通过烟雾拖延进攻，这些贡献在普通胜负标签中并不会单独出现。",
    57: " 角色分类还会影响玩家数据的解释方式：同样的死亡次数，在突破位上可能来自主动进攻，在哨位上则可能说明防守位置被突破。",
    58: " 因此，地图因子并不是为了人为提高某张地图的权重，而是把地图结构对角色发挥和玩家熟练度的影响纳入评分解释。",
    60: " 这种写法牺牲了一部分工程扩展性，但部署门槛低，打开后端服务和浏览器页面即可完成主要功能演示。",
    68: " 这也是优化ELO从K值入手的原因：不改变ELO胜率期望公式的主体，只让评分更新幅度随对局信息发生小范围变化。",
    69: " 通过这种补正，系统不会把所有角色都放到同一KDA模板下评价，而是在保留胜负结果的基础上调整不同职责的分数波动。",
    88: " 这些字段与后文特征工程一一对应，便于从数据生成、模型训练到页面展示形成闭环，而不是只停留在公式推导层面。",
    103: " 这部分实验也解释了为什么系统最终采用GBDT作为主要预测模型：它在保持可解释特征输入的同时，能够处理多因素之间的非线性关系。",
    108: " 因此，消融实验不仅验证模型效果，也反过来说明第三章三个修正因子的设计具有必要性，而不是只为了增加公式复杂度。",
    126: " 这些模板还承担演示作用，用户在不了解公式细节时，也能通过输入场景观察两种ELO算法在加分和扣分上的差异。",
    131: " 该接口减少了前端多次请求，也让历史记录页面能够按统一格式展示不同业务来源的数据，方便用户回溯自己的操作。",
    137: " 排名数据不是单纯按用户名排序，而是以最新评分为核心，再补充胜率和KDA，使排行榜既反映等级分，也保留玩家近期表现线索。",
    138: " 领奖台用于突出头部玩家，表格用于完整查询；两种视图使用同一接口数据，减少了前后端之间重复维护排名逻辑的成本。",
    143: " 这种设计借鉴了竞技游戏常见的段位表达方式，既保留ELO连续分值，又避免用户只看到一个抽象数字。",
    165: " 对毕业设计而言，这种集中式前端结构便于检查页面逻辑，也便于在没有复杂构建环境的电脑上快速运行。",
    166: " Toast用于保存、删除、计算等操作后的反馈，段位弹窗用于用户自定义展示身份，图表刷新则保证数据变化后页面结果同步更新。",
}

SECOND_EXTENSIONS: dict[int, str] = {
    28: " 这样处理后，文献综述与后文算法设计之间形成对应关系：前者指出标准ELO信息不足，后者再说明KDA、角色和地图三类修正因子从何而来。",
    29: " 这一段的重点不是罗列模型名称，而是说明胜负预测研究已经从单一统计指标转向多特征建模，本文的模型比较也沿着这条思路展开。",
    32: " 因此，论文在研究空白处强调的是场景适配问题，而不是简单宣称某一种模型优于其他模型。",
    39: " 这种低成本更新机制也是本文保留ELO框架的原因：它适合持续记录玩家战绩，并能与Web系统中的历史记录模块结合。",
    48: " 通过保留逻辑回归，实验结果可以更清楚地显示非线性模型带来的增益，而不是只在复杂模型之间比较高低。",
    50: " 随机森林的结果还可以帮助判断：若它与GBDT差距较小，说明树模型整体有效；若差距较大，则说明Boosting的逐步修正更适合当前数据。",
    54: " 这四个指标也便于在表格中横向比较模型，读者可以同时看到总体正确率和对胜利样本识别能力的变化。",
    56: " 这些局内因素虽然不一定都能直接量化，但它们解释了为什么本文至少要引入KDA、角色和地图三个可获取变量。",
    58: " 玩家对某张地图的熟悉程度也可能通过站位、技能释放时机和转点路线体现，因此地图变量对胜负预测具有间接价值。",
    60: " 这一选择与系统规模相匹配：项目需要展示算法、预测、历史记录和可视化，而不是构建大型前端工程。",
    68: " 如果不处理这一点，系统只能告诉用户分数变了多少，却无法解释为什么同样胜利时不同玩家应当获得不同评分变化。",
    88: " 由于所有字段都由脚本生成，实验还具备可重复性，便于在同一数据分布下比较不同模型和消融设置。",
    103: " 该结论与后续系统实现相连：预测接口最终调用表现较好的模型，而不是在页面中只展示静态公式。",
    108: " 这也说明，本系统中的算法优化和模型实验不是两个孤立部分，前者提供领域修正，后者检验这些修正是否能转化为预测收益。",
    131: " 这种聚合接口的写法也降低了前端状态管理复杂度，页面只需处理一组返回对象即可渲染三个历史列表。",
    165: " 当然，这种组织方式也意味着后续若继续扩大功能，仍需要拆分组件和引入构建工具；论文中将其限定为毕业设计规模的实现方案。",
}


def count_han(document: Document) -> int:
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return len(HAN_RE.findall("\n".join(texts)))


def is_green(paragraph) -> bool:
    for run in paragraph.runs:
        try:
            color = run.font.color.rgb if run.font.color and run.font.color.rgb else None
        except Exception:
            color = None
        if str(color).upper() == "1F8F3A":
            return True
    return False


def replace_green(paragraph, text: str) -> None:
    first_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = GREEN
    if first_run is not None:
        run.bold = first_run.bold
        run.italic = first_run.italic
        run.underline = first_run.underline
        run.font.size = first_run.font.size
        run.font.name = first_run.font.name


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = WORK_DIR / f"round2_before_length_restore_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup)
    document = Document(str(DOCX_PATH))
    before = count_han(document)

    changed = 0
    for index, extension in EXTENSIONS.items():
        if index < len(document.paragraphs) and is_green(document.paragraphs[index]):
            base = document.paragraphs[index].text.rstrip()
            if extension.strip() not in base:
                replace_green(document.paragraphs[index], LETTER_SPACE_RE.sub("", base + extension))
                changed += 1
    for index, extension in SECOND_EXTENSIONS.items():
        if index < len(document.paragraphs) and is_green(document.paragraphs[index]):
            base = document.paragraphs[index].text.rstrip()
            if extension.strip() not in base:
                replace_green(document.paragraphs[index], LETTER_SPACE_RE.sub("", base + extension))
                changed += 1

    after = count_han(document)
    temp = WORK_DIR / f"round2_length_restored_{timestamp}.docx"
    document.save(temp)
    for target in (DOCX_PATH, FINAL_DOCX):
        if target.exists():
            target.chmod(0o666)
        shutil.copyfile(temp, target)
    temp.unlink(missing_ok=True)

    old = REPORT_PATH.read_text(encoding="utf-8") if REPORT_PATH.exists() else "# AIGC 改稿说明\n"
    REPORT_PATH.write_text(
        old.rstrip()
        + "\n\n## 长度回补\n\n"
        + f"- 回补前汉字数：{before}\n"
        + f"- 回补后汉字数：{after}\n"
        + f"- 回补绿色热点段：{changed}\n"
        + "- 回补内容只来自原稿已有的算法、实验、接口和页面说明，不新增报告外事实。\n",
        encoding="utf-8",
    )

    print(f"backup={backup}")
    print(f"changed={changed}")
    print(f"before_han={before}")
    print(f"after_han={after}")


if __name__ == "__main__":
    main()
