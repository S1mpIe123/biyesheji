from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT = BASE_DIR / "毕业答辩背诵提纲-张博源.docx"


def set_font(run, size=11, bold=False):
    """统一设置中英文字体，避免 Word 打开后出现字体混乱。"""
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    """添加章节标题，保留 Word 标题层级，方便自动导航。"""
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    set_font(run, size=16 if level == 1 else 13, bold=True)
    return paragraph


def add_paragraph(doc, text, size=11, bold_prefix=None):
    """添加正文段落；可把问答中的问题前缀加粗，便于背诵时快速定位。"""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(22)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        run1 = paragraph.add_run(bold_prefix)
        set_font(run1, size=size, bold=True)
        run2 = paragraph.add_run(text[len(bold_prefix):])
        set_font(run2, size=size)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=size)
    return paragraph


def add_bullet(doc, text):
    """添加简短记忆点，不用于论文正文，仅用于答辩背诵提纲。"""
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_font(run, size=11)


def build_document():
    """生成答辩背诵文档主体。"""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("毕业答辩背诵提纲")
    set_font(title_run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("基于 Spring Boot 和 Vue 的农产品产销对接平台设计与实现")
    set_font(subtitle_run, size=12)

    add_heading(doc, "一、30秒项目总述", 1)
    add_paragraph(
        doc,
        "我的毕业设计是基于 Spring Boot 和 Vue 的农产品产销对接平台。系统不是单纯的农产品商城，而是在商城基础上扩展了供应发布、采购需求、智能匹配推荐、撮合记录、订单库存同步和数据分析等功能。后端采用 Spring Boot、MyBatis 和 MySQL，前端采用 Vue、Element Plus 和 Axios。项目的核心是通过规则权重算法计算供应信息和采购需求的匹配度，并通过撮合记录和库存同步机制保证供需对接过程可追踪、库存变化可控制。系统完成了管理员、供应商、采购商和普通购买用户四类角色的主要业务流程，基本实现了农产品供应方、采购方和普通购买用户之间的信息对接与交易闭环。",
    )

    add_heading(doc, "二、项目亮点", 1)
    add_bullet(doc, "项目定位不是普通商城，而是“商城交易 + 供需撮合”的产销对接平台。")
    add_bullet(doc, "系统包含管理员、供应商、采购商、普通购买用户四类角色，业务分工清晰。")
    add_bullet(doc, "实现了基于规则权重的供需匹配推荐，能按匹配度排序推荐供应或采购需求。")
    add_bullet(doc, "供应审核通过后可同步为商城商品，商城下单会同步扣减来源供应库存。")
    add_bullet(doc, "撮合记录具有状态流转，确认、取消、完成等状态会联动库存锁定和释放。")
    add_bullet(doc, "数据分析看板可展示供应量、采购需求、撮合成功率、订单闭环率和库存风险。")

    add_heading(doc, "三、高频问题与回答", 1)
    qa_items = [
        (
            "1. 你的项目主要解决什么问题？",
            "本系统主要解决农产品流通过程中供需信息分散、采购需求和供应库存难以及时匹配的问题。传统商城更偏向消费者购买商品，而我的系统在商城基础上增加了供应商发布供应、采购商发布需求、系统智能推荐、双方撮合、订单和库存同步等功能，使供应方、采购方和普通消费者都能在同一平台完成信息对接和交易管理。",
        ),
        (
            "2. 你的系统和普通商城有什么区别？",
            "普通商城主要是商品展示和下单购买，我的系统多了一条产销对接业务链。供应商可以发布供应信息，采购商可以发布采购需求，系统根据产品名称、分类、价格、地区、时间、数量等条件计算匹配度，生成推荐结果。撮合成功后会形成撮合记录，并联动库存变化，所以它是商城交易和供需撮合结合的平台。",
        ),
        (
            "3. 项目用了什么算法？",
            "项目使用的是基于规则权重的供需匹配算法，不是机器学习算法。算法先判断农产品名称是否匹配，如果名称完全不匹配，直接过滤，避免只因为同分类就出现错误推荐。名称匹配后，再从分类、时间、价格、地区、数量等维度累加得分，最后按照得分排序，超过阈值的记录才进入推荐列表。",
        ),
        (
            "4. 匹配算法的权重怎么设计？",
            "名称匹配是核心条件，完全一致 50 分，包含关系 40 分，有一定共同字符 25 分；分类一致 10 分；时间在 7 天内 20 分，30 天内 10 分；供应价格低于期望价格 15 分，超出 20% 以内 8 分；地区匹配 15 分，否则给少量基础分；数量满足需求 10 分，不完全满足 5 分。总分最高限制为 100 分，系统筛选 60 分以上的结果。",
        ),
        (
            "5. 为什么不用机器学习算法？",
            "因为本项目真实历史成交数据量有限，直接训练机器学习模型缺少可靠数据基础。规则权重算法更适合当前毕业设计场景，优点是实现稳定、逻辑清楚、结果可解释，答辩时也能说明每个推荐结果为什么出现。后续如果平台积累了足够成交数据，可以扩展为协同过滤、学习排序或机器学习推荐模型。",
        ),
        (
            "6. 库存一致性怎么保证？",
            "系统把供应信息的剩余数量作为源头库存。供应审核通过后可以同步生成商城商品，用户在商城下单时扣减商品库存，如果该商品来源于供应信息，也同步扣减供应剩余量。订单取消时会返还商品库存和供应库存。撮合记录从待确认进入已确认、已支付、待发货、待收货、已完成等状态时，会锁定对应供应量和采购需求量；未完成前取消或退回，则释放库存。后端这些更新方法使用事务控制，避免只改订单不改库存的情况。",
        ),
        (
            "7. 撮合记录状态怎么设计？",
            "撮合记录用于记录供应商和采购商之间的一次对接过程。初始状态是待确认，确认后进入履约流程，后续可以流转到已支付、待发货、待收货、已完成、售后中等状态。状态变化不仅是页面展示，还会触发业务逻辑，例如确认后锁定库存，取消时释放库存，已完成后不再回补库存，避免历史成交影响剩余库存。",
        ),
        (
            "8. 权限怎么控制？",
            "系统按角色区分功能入口。管理员负责用户、分类、商品、审核、订单和数据分析；供应商主要发布供应信息、查看推荐采购需求、处理发货；采购商发布采购需求、浏览供应信息、发起采购意向；普通购买用户主要使用商城购买功能。部分关键规则放在后端业务层，例如供应商不能进行普通购买，普通用户只能查看自己的订单，发货动作也会判断操作者是否是对应供货方。",
        ),
        (
            "9. 数据分析看板有什么作用？",
            "数据分析看板主要服务管理员运营管理，可以查看供应信息数量、采购需求数量、撮合成功率、订单闭环率、产地分布、品类结构和库存风险等指标。它的作用不是单纯展示图表，而是帮助管理员了解平台供需是否平衡、哪些品类更活跃、哪些商品存在库存风险，为后续审核和运营调整提供参考。",
        ),
        (
            "10. 项目有哪些不足？",
            "目前系统的登录认证相对简单，没有完整引入 Spring Security 和 JWT；智能匹配是规则权重算法，还不是基于大量历史数据训练的模型；数据分析中的 AI 分析更偏本地规则生成，没有接入真实大模型接口；并发库存控制主要依赖事务，暂未引入 Redis 分布式锁。这些属于毕业设计阶段的边界，后续可以继续扩展。",
        ),
    ]

    for question, answer in qa_items:
        add_heading(doc, question, 2)
        add_paragraph(doc, answer)

    add_heading(doc, "四、演示前检查清单", 1)
    checklist = [
        "准备四类账号：管理员、供应商、采购商、普通购买用户。",
        "准备一条完整演示链路：发布供应、发布需求、管理员审核、智能推荐、生成撮合记录、状态流转、商城下单、库存同步变化。",
        "提前检查数据库演示数据，保证有供应、有采购、有推荐结果、有订单、有分析看板数据。",
        "背熟算法关键词：名称优先、规则权重、多维评分、阈值过滤、按分排序。",
        "主动准备不足说明：认证体系、推荐算法、真实 AI 接口、并发库存控制都可以作为后续优化方向。",
    ]
    for item in checklist:
        add_bullet(doc, item)

    add_heading(doc, "五、最短背诵版", 1)
    add_paragraph(
        doc,
        "本项目是一个农产品产销对接平台，不是单纯商城。它围绕供应商、采购商、普通购买用户和管理员四类角色，实现了供应发布、采购需求、智能匹配、撮合记录、订单库存同步和数据分析。核心算法是基于规则权重的供需匹配算法，先按农产品名称过滤，再按分类、时间、价格、地区和数量累加得分，筛选高分结果进行推荐。系统亮点在于把供需撮合和商城购买打通，供应审核通过后可同步为商品，商城下单和撮合确认都会联动库存变化，从而形成信息对接、交易管理和库存追踪的闭环。",
    )

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    # 保存后检查 docx 的 zip 包结构，能降低生成损坏文件的风险。
    with zipfile.ZipFile(output) as archive:
        bad_file = archive.testzip()
    if bad_file:
        raise RuntimeError(f"DOCX 文件结构异常: {bad_file}")
    print(output)
