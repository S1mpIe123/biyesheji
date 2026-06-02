from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor


# 二次返工脚本：
# 1. 只处理文本，不重建图片、表格和题注结构；
# 2. 删除中文与英文/数字之间的多余空格；
# 3. 对绿色 AIGC 热点段落做保守压缩，保留项目事实。
WORK_DIR = Path("output/aigc_revision")
WORKING_DOCX = WORK_DIR / "working_aigc_report.docx"
FINAL_DOCX = WORK_DIR / "免费_Word标红版_AIGC检测报告_[基于springboot+vue的农产品]_改稿版.docx"
REPORT_MD = WORK_DIR / "AIGC改稿说明.md"
GREEN = RGBColor(0x1F, 0x8F, 0x3A)

HAN_RE = re.compile(r"[\u4e00-\u9fff]")
CJK_ASCII_SPACE_RE = re.compile(
    r"(?<=[\u4e00-\u9fff])\s+(?=[A-Za-z0-9])|(?<=[A-Za-z0-9])\s+(?=[\u4e00-\u9fff])"
)


# 这些替换只用于绿色段落，目标是压缩机械化表达，不改变系统事实。
GREEN_COMPACTION_RULES: list[tuple[str, str]] = [
    ("本系统", "系统"),
    ("本课题", "课题"),
    ("本项目", "项目"),
    ("本文", "文中"),
    ("当前", ""),
    ("主要", ""),
    ("已经", "已"),
    ("能够", "能"),
    ("可以", "可"),
    ("需要", "需"),
    ("进行", ""),
    ("完成", "实现"),
    ("统一", "集中"),
    ("后续可", "后续可"),
    ("后续可以", "后续可"),
    ("这样处理后，", "这样，"),
    ("也就是说，", ""),
    ("而不是", "而非"),
    ("并且", "并"),
    ("同时", "并"),
    ("由于", "因"),
    ("如果", "若"),
    ("已经", "已"),
    ("反复", "多次"),
    ("一条", "一项"),
]


# 少数较长热点段采用人工压缩，避免纯规则删除造成语义断裂。
TARGETED_REWRITES: dict[int, str] = {
    22: "系统按前后端分离实现。后端使用Spring Boot3.3.1、MyBatis和MySQL，保存用户、分类、商品、供应、采购需求、撮合记录和订单等数据；前端使用Vue3.3.4、Element Plus和Axios实现页面与接口交互。四类角色分别承担审核维护、货源发布、需求发布和商品购买。匹配推荐在MatchRecommendService中按名称、分类、时间、价格、地区和数量评分；撮合记录和订单再联动库存变化，形成信息发布、推荐、确认、下单和追踪链路。",
    119: "近几年农业数字化持续推进，农产品流通也更多依赖线上信息平台。政策文件提出推进数字乡村建设，相关研究从物联网、大数据和供需匹配角度讨论农业产销协同[1][2]。在这一背景下，课题用Spring Boot和Vue实现供应发布、采购需求、推荐撮合和订单处理等功能。",
    120: "农产品流通中常见的问题是信息到达不及时。供应商掌握产地、数量和上市时间，却难找到合适采购方；采购商知道需求，也常在多个渠道询价。已有研究从信息不对称、小农经营分散和流通层级较多等角度解释这一现象[3]。系统将其转化为供应信息、采购需求和撮合记录等数据对象。",
    126: "从流通效率看，多层中间渠道会增加沟通和议价成本。平台把货源和需求放在同一系统中，支持按产品、分类、地区、价格等条件筛选。郭杰等研究说明，合理匹配机制有助于提高供需对接效率[5]。本系统采用规则评分，先筛出较接近的组合，减少用户逐条查找的工作量。",
    121: "传统交易依赖批发市场、经纪人、产地集散点和零散网络发布。它们能完成交易，但信息分散，后续查询不便。董秋丽等指出，供需信息平台仍存在功能单一、更新不及时、匹配能力不足等问题[4]。因此，系统把审核、推荐、撮合状态和库存变化放入同一流程。",
    131: "功能设计上，系统在农产品商城外增加产销对接链路。供应审核通过后可同步为商品；普通用户下单时扣减商品库存和来源供应数量；撮合确认后锁定供应量和采购需求量。商城与供需大厅通过source_supply_id、quantity等字段关联，销售行为能回写供应端数据。",
    210: "从使用需求看，供应商需要展示货源，采购商需要快速找货，普通用户需要购买入口，管理员需要审核和运营数据。杨光梅指出，数字化平台能缓解小规模主体信息渠道有限的问题[11]。系统据此拆分为供应发布、需求发布、推荐撮合、订单和数据看板等功能。",
    127: "从信息整合看，同类农产品可能在某地供应充足，在另一地却难找货源。系统集中保存供应中的产品、产地、价格、数量、上市时间，以及需求中的数量、期望价、收货地区和截止时间，便于查询比较，也减少多渠道核对成本。",
    140: "需求整理阶段按管理员、供应商、采购商和普通购买用户划分。管理员负责审核、分类维护和数据查看；供应商发布供应并处理直供订单；采购商发布需求、查看推荐并发起意向；普通用户浏览商品并下单。接口和性能要求围绕这些页面展开。",
    641: "匹配推荐测试使用同一组红富士苹果数据。需求为3000斤、期望价5.5元，供应为5000斤、供应价5.0元。算法计算名称50分、分类10分、价格15分、地区5分、数量10分，时间差超过30天不得分，合计90分。页面首条推荐与计算一致，发起2000斤采购意向后生成待确认撮合记录。图6-3展示测试结果。",
    302: "推荐计算读取已通过且剩余量大于零的供应和采购需求，再逐对比较。算法先处理产品名称，去掉“求购”“供应”“采购”“新鲜”等干扰词，并识别苹果、玉米、猕猴桃及红富士、甜玉米等产品或品种词。只有名称一致、包含或核心品类相近的组合才继续评分，避免仅因分类相同产生误推荐。",
    145: "测试验证阶段在本地检查主要业务路径，包括注册登录、供应提交与审核、采购需求审核、推荐得分、撮合状态流转、商城下单扣库和订单取消返还。测试目标是确认常见操作能按规则更新状态和数量，而非给出性能基准。图1-1展示研究方法与技术路线。",
    364: "SupplyService.add插入供应前先处理关键字段。单位为“吨”时，quantity和totalQuantity换算为斤，unit保存为斤；status为空则设为待审核，createTime为空则写当前时间；若同一用户已发布同品种采购需求，系统拒绝再次发布供应，避免自买自卖式数据。供应状态为已通过时，再同步生成或更新商城商品。",
    215: "页面交互需保持基本流畅。商品列表、供应大厅和采购需求大厅采用分页查询，单页数据量不宜过大；图片上传完成后应返回可访问地址；路由切换和表格刷新不应长时间空白。因课题未做专项压力测试，文中只提出响应及时性的设计要求，不写无法验证的性能承诺。",
    633: "供应发布与审核测试中，供应商test_farmer提交“红富士苹果”货源，包含分类、产地、单位、价格、数量和上市时间。记录先进入待审核状态；管理员通过后，状态变为已通过，并同步生成source_supply_id指向该供应的商城商品，库存与供应剩余量一致，结果符合预期。",
}


def iter_paragraphs(document: Document):
    """遍历正文段落和表格单元格段落，避免漏掉表格中的文字。"""
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def is_green_paragraph(paragraph) -> bool:
    """判断段落是否为上轮 AIGC 改写后标记的绿色段落。"""
    for run in paragraph.runs:
        try:
            color = run.font.color.rgb if run.font.color and run.font.color.rgb else None
        except Exception:
            color = None
        if str(color).upper() == "1F8F3A":
            return True
    return False


def replace_paragraph_text(paragraph, text: str) -> None:
    """替换整段文字并保持段落级样式；绿色用于标识本轮仍属于热点处理区。"""
    first_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = GREEN
    if first_run is not None:
        run.bold = first_run.bold
        run.italic = first_run.italic
        run.underline = first_run.underline
        run.font.size = first_run.font.size
        run.font.name = first_run.font.name


def compact_green_text(text: str) -> str:
    """对绿色段落做保守压缩，减少模板化词语和无效连接词。"""
    compacted = text
    for old, new in GREEN_COMPACTION_RULES:
        compacted = compacted.replace(old, new)
    compacted = re.sub(r"。(?=图\d+-\d+)", "，", compacted)
    compacted = re.sub(r"\s+", " ", compacted).strip()
    return compacted


def normalize_cjk_ascii_spacing_in_text_nodes(paragraph) -> int:
    """清理一个段落内中文与英文/数字之间的空格，跨 run 边界也处理。"""
    text_nodes = paragraph._p.xpath(".//w:t")
    changes = 0

    for node in text_nodes:
        old = node.text or ""
        new = CJK_ASCII_SPACE_RE.sub("", old)
        if new != old:
            node.text = new
            changes += 1

    for left, right in zip(text_nodes, text_nodes[1:]):
        left_text = left.text or ""
        right_text = right.text or ""
        new_left = re.sub(r"(?<=[\u4e00-\u9fff])\s+$", "", left_text) if re.match(r"^[A-Za-z0-9]", right_text) else left_text
        new_left = re.sub(r"(?<=[A-Za-z0-9])\s+$", "", new_left) if re.match(r"^[\u4e00-\u9fff]", right_text) else new_left
        new_right = re.sub(r"^\s+(?=[A-Za-z0-9])", "", right_text) if re.search(r"[\u4e00-\u9fff]$", new_left) else right_text
        new_right = re.sub(r"^\s+(?=[\u4e00-\u9fff])", "", new_right) if re.search(r"[A-Za-z0-9]$", new_left) else new_right
        if new_left != left_text:
            left.text = new_left
            changes += 1
        if new_right != right_text:
            right.text = new_right
            changes += 1

    return changes


def count_document(document: Document) -> dict[str, int]:
    """输出验证用统计：中文字符、绿色段落和中英相邻空格数量。"""
    texts: list[str] = []
    green_count = 0
    green_han = 0
    for paragraph in iter_paragraphs(document):
        text = paragraph.text
        texts.append(text)
        if is_green_paragraph(paragraph):
            green_count += 1
            green_han += len(HAN_RE.findall(text))
    all_text = "\n".join(texts)
    return {
        "total_han_chars": len(HAN_RE.findall(all_text)),
        "cn_en_space_matches": len(CJK_ASCII_SPACE_RE.findall(all_text)),
        "green_count": green_count,
        "green_han_chars": green_han,
    }


def append_report(before: dict[str, int], after: dict[str, int], backup_path: Path) -> None:
    """把二次返工记录写入单独说明文件，不污染论文正文。"""
    addition = [
        "",
        "## 2026-05-11 二次返工：中英空格与字数控制",
        "",
        f"- 备份稿：{backup_path.resolve()}",
        f"- 中文字符数（不含英文、标点）：{before['total_han_chars']} -> {after['total_han_chars']}",
        f"- 绿色热点段中文字符数：{before['green_han_chars']} -> {after['green_han_chars']}",
        f"- 中文与英文/数字之间的空格匹配数：{before['cn_en_space_matches']} -> {after['cn_en_space_matches']}",
        "- 本轮只压缩绿色 AIGC 热点段落，并清理全文中英相邻空格；图表、题注、表格结构未重建。",
    ]
    old = REPORT_MD.read_text(encoding="utf-8") if REPORT_MD.exists() else "# AIGC 改稿说明\n"
    REPORT_MD.write_text(old.rstrip() + "\n" + "\n".join(addition) + "\n", encoding="utf-8")


def main() -> None:
    if not WORKING_DOCX.exists():
        raise FileNotFoundError(WORKING_DOCX)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = WORK_DIR / f"working_aigc_report_spacing_backup_{timestamp}.docx"
    shutil.copy2(WORKING_DOCX, backup_path)

    document = Document(str(WORKING_DOCX))
    before = count_document(document)

    green_index = 0
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if is_green_paragraph(paragraph):
            if paragraph_index in TARGETED_REWRITES:
                replace_paragraph_text(paragraph, TARGETED_REWRITES[paragraph_index])
            else:
                replace_paragraph_text(paragraph, compact_green_text(paragraph.text))
            green_index += 1

    spacing_changes = 0
    for paragraph in iter_paragraphs(document):
        spacing_changes += normalize_cjk_ascii_spacing_in_text_nodes(paragraph)

    after = count_document(document)
    document.save(WORKING_DOCX)
    shutil.copy2(WORKING_DOCX, FINAL_DOCX)
    append_report(before, after, backup_path)

    print(f"backup={backup_path}")
    print(f"spacing_node_changes={spacing_changes}")
    print(f"before={before}")
    print(f"after={after}")
    print(f"final={FINAL_DOCX}")


if __name__ == "__main__":
    main()
