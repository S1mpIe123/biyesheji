from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


# VALORANT/ELO 检测报告改写脚本。
# 只替换被 AIGC 颜色标出的正文段落，不重建图片、公式、表格和题注。
WORK_DIR = Path("output/aigc_revision_valorant")
DOCX_PATH = WORK_DIR / "working_valorant_aigc_report.docx"
FINAL_DOCX = WORK_DIR / "免费_Word标红版_AIGC检测报告_[基于ELO机制的VALORANT玩家 胜]_改稿版.docx"
REPORT_PATH = WORK_DIR / "AIGC改稿说明.md"
GREEN = RGBColor(0x1F, 0x8F, 0x3A)

HAN_RE = re.compile(r"[\u4e00-\u9fff]")
LETTER_SPACE_RE = re.compile(r"(?<=[\u4e00-\u9fff]) +(?=[A-Za-z])|(?<=[A-Za-z]) +(?=[\u4e00-\u9fff])")
RISK_COLORS = {"F12828", "F39800", "9D91E9"}


# key 为 python-docx 读取到的段落序号，value 为保留原事实后的降 AIGC 改写文本。
REWRITES: dict[int, str] = {
    28: "国内研究对ELO的使用多见于游戏匹配、玩家评价和竞技系统评测。部分研究把ELO框架放到MOBA游戏的匹配质量讨论中，关注K因子调整对玩家体验的影响[7]。徐英卓、郭博、王六鹏（2023）使用GBDT处理游戏销量预测，说明梯度提升树在游戏数据分析中有一定适用性。相比只依据胜负更新分值的传统ELO机制，VALORANT场景还需要把玩家表现、角色定位和地图差异纳入评分解释。",
    29: "游戏胜负预测是竞技数据分析中的常见问题，任务目标是在赛前或赛中估计比赛结果概率。现有方法大致包括时序模型、图神经网络和集成学习。以MOBA为对象的研究较多，例如将经济差、经验差、推塔数等时间序列特征输入Seq2Seq模型，用来描述局势变化。此类方法提供了思路，但其特征来源与VALORANT的回合制战术射击机制并不完全一致。",
    30: "图神经网络方向的研究通常把玩家看作节点，把击杀、助攻或协作关系看作边。Sun等人（2021）用GNN建模MOBA团队内部的互动结构，并在Dota2和英雄联盟数据集上取得较好结果[10]。Li等人（2022）则采用集成学习融合多个基础模型，在DOTA2数据集上报告了较高准确率[11]。这些工作证明复杂关系建模有价值，但也依赖与具体游戏机制匹配的数据结构。",
    31: "梯度提升决策树在游戏业务预测中也有应用。刘镇恺（2024）围绕游戏用户流失预测比较多种机器学习算法，并在GBDT基础上结合业务特征进行工程处理和参数优化[12]。该研究说明，GBDT适合处理含有多类行为特征的预测任务。本文借鉴这一思路，把ELO、KDA、角色和地图等变量整理成结构化特征后用于胜负预测。",
    32: "上述研究各有进展，但直接迁移到VALORANT仍存在限制。VALORANT具有5v5战术射击、回合制攻守轮换、角色技能和枪法共同作用等特点，MOBA中的经济曲线、兵线和团队图结构不能简单替换使用。市场上部分预测工具仍主要依赖ELO差值作线性判断，特征工程和模型解释深度不足。本文因此围绕VALORANT机制重新组织评分因子和预测特征。",
    34: "数据可视化是电竞数据分析系统面向用户的呈现层。它把模型计算结果转为玩家更容易理解的图表和指标。郑涛、陈婷婷等人（2022）基于Flask和ECharts实现英雄联盟LPL赛事数据可视化，展示KDA、英雄禁用率和经济走势等内容[14]。该工作对本文的启发在于：算法结果需要通过仪表盘、排行榜和趋势图等形式进入用户操作流程。",
    35: "电竞产业研究为系统需求提供了背景。游继之和布特（2018）分析了我国电子竞技产业链结构和发展趋势[15]，说明围绕玩家数据、赛事表现和竞技评价的工具具有现实需求。本文的系统设计也落在这一类数据服务场景中。",
    39: "ELO评分由Arpad Elo提出，最初用于国际象棋等级分管理，后被USCF和FIDE采用[1]。其基本思想是用一个连续分值近似表示竞技者水平，再将双方分差映射为胜率期望。比赛结束后，系统根据实际结果与期望结果的差距更新分值。这个机制计算简单，适合连续比赛记录，但它默认比赛结果能充分代表玩家表现。",
    41: "标准ELO的优点是公式简洁、参数容易解释，且每场比赛只需常量级计算即可更新分数。问题也很明显：K因子通常固定，期望得分主要依赖评分差，实际结果被压缩为胜负二值。棋类一对一场景中这种处理可以接受，但VALORANT属于团队战术射击，玩家KDA、角色职责、地图结构都会影响贡献，单靠胜负更新容易丢失信息。",
    48: "逻辑回归是常用的对数线性分类模型。它先计算特征向量与权重的线性组合，再通过Sigmoid函数输出0到1之间的概率值。模型参数少、训练快、结果容易解释，适合作为胜负预测的基线模型。本文将其与随机森林、GBDT放在同一数据集上比较，用来观察复杂模型是否真正带来提升。",
    50: "随机森林是Bagging思想下的集成模型。训练时，算法从原始样本中多次自助采样，分别训练多棵CART决策树；每个分裂节点只在随机特征子集上寻找划分。预测时，多棵树通过投票给出结果。随机森林能够降低单棵树过拟合风险，并能处理非线性特征关系，因此适合作为GBDT之外的对比模型。",
    54: "模型评估使用准确率、精确率、召回率和F1分数。准确率反映测试集整体预测正确比例；精确率关注预测为胜利的样本中有多少确实胜利；召回率关注实际胜利样本中有多少被识别出来；F1分数综合精确率和召回率。四项指标合用，可以避免只看准确率而忽略类别识别偏差。",
    56: "VALORANT采用5v5攻守轮换模式，常规对局最多24回合，先取得13回合的一方获胜。进攻方需要安装并保护Spike，防守方则阻止安装或在安装后完成拆除。游戏在FPS枪法基础上加入角色技能，使胜负受到枪法、站位、道具配合、地图理解等因素共同影响。这也是本文不直接沿用普通ELO的原因。",
    57: "角色按战术职能分为突破位、发起位、控场位和哨位。突破位负责打开交火空间，通常击杀和死亡都较高；发起位依靠侦察、闪光等技能帮助团队获取信息；控场位用烟雾和区域控制改变交战条件；哨位负责防守、断后和信息保护。不同职责会影响KDA分布，因此评分时需要避免把所有角色放在同一表现标准下比较。",
    58: "不同地图会改变角色价值和对局节奏。Ascent强调中路控制，Bind有传送门带来的快速转点，Haven具有三个据点，Split和Icebox更突出垂直空间，Breeze和Pearl等地图则影响远距离交火和区域控制。地图不是单纯背景，而是影响角色选择、交战方式和胜率判断的结构变量。",
    60: "前端采用Vue.js 2.7。该框架通过响应式数据绑定和模板语法把数据变化映射到界面更新，减少手动DOM操作。系统使用CDN运行时，不引入复杂构建工具链，适合将页面逻辑集中在index.html中。计算属性、条件渲染和事件绑定用于支撑多标签页切换、表单输入、图表刷新和用户反馈。",
    62: "login.html使用Canvas 2D API实现动态粒子背景。页面设置多个画布层，分别绘制慢速背景粒子、曲线运动粒子和闪烁粒子，通过requestAnimationFrame持续更新。该设计主要服务于电竞风格的视觉表现，不影响用户认证逻辑。粒子越界后重置位置和运动参数，以保持背景动画连续。",
    68: "标准ELO的第一项限制是K因子固定。无论比赛是大比分碾压还是艰难取胜，也无论玩家是核心Carry还是被队友带动，评分变化都主要由赛前分差和胜负结果决定。这样的处理会把KDA、回合贡献和比赛过程压缩掉。对VALORANT这类团队游戏而言，固定K值难以反映单局表现差异。",
    69: "第二项限制来自角色差异。突破位和哨位即使ELO相同，合理表现也可能完全不同：前者承担进攻风险，击杀和死亡都偏高；后者更强调生存、控区和信息保护。若评分只看胜负，突破位的高风险贡献和哨位的稳定防守都难被准确区分。因此本文在K因子中加入角色权重，用于缓解职责差异造成的评价不对称。",
    70: "第三项限制是地图效应被平均化。VALORANT地图在据点数量、转点方式、垂直结构和技能价值上差异明显。以Icebox为例，复杂的立体结构会放大发起位侦察技能的作用；Ascent则更强调中路控制和控场配合。若所有地图结果以同一方式进入ELO更新，玩家的地图适应性就无法体现。",
    73: "KDA修正因子用于把击杀、死亡和助攻转化为K因子的调节系数。公式先计算kda_raw=(kills+assists)/max(deaths,1)，用max(deaths,1)处理零死亡边界。随后对极端KDA进行裁剪和映射，避免少数异常数据造成评分剧烈波动。这样既保留个人表现信息，也控制其对评分更新的影响范围。",
    77: "角色权重区间设置为[0.97,1.05]，幅度较窄。这样做是为了让角色差异起到补正作用，而不是取代KDA或胜负结果。FPS游戏中枪法和临场表现仍是主要因素，角色权重只用于修正不同职责下表现指标的天然差异，避免评分因角色系数过大而漂移。",
    81: "三个修正因子组合后形成动态K因子。综合调节系数K_adjust由kda_factor、role_factor和map_factor相乘得到。该形式不再额外设置主观权重，减少调参负担，也便于解释每个因子对评分更新幅度的影响。最终评分变化仍保留标准ELO的基本结构。",
    88: "训练数据由generate_data.py按预设分布模拟生成，共1500条对局记录。数据设计考虑样本量、特征覆盖和胜负比例平衡：样本规模需支撑训练与评估，特征需覆盖ELO、KDA、角色和地图等维度，胜负比例控制在接近均衡的范围。每条记录包含玩家ELO、对手ELO、KDA、角色、地图等字段。",
    89: "1500条样本是在训练充分性和过拟合风险之间的折中。数据集共有12个特征，样本量相对特征数较充足，可支持GBDT中多棵树的训练。实验按4:1划分训练集和测试集，其中1200条用于模型拟合和参数搜索，300条用于最终评估。",
    91: "特征工程围绕四个维度展开，共设计12个核心特征。选择原则是覆盖影响胜负的主要因素，同时避免特征数量过多。这样可以在样本量有限的情况下减少过拟合风险，并使模型输入与VALORANT对局机制保持对应。",
    92: "ELO维度包含player_elo、opponent_elo、elo_diff和elo_pred。前两个字段记录双方评分，elo_diff表示分差方向和幅度，elo_pred则把第三章优化ELO算法得到的预测结果作为级联特征输入机器学习模型，使模型能够利用ELO算法提供的先验信息。",
    93: "KDA维度包含kda_ratio、kda_factor和kda_bucket。kda_ratio保留连续比值，kda_factor直接使用前文设计的KDA修正系数，kda_bucket把表现划分为低、中、高三个区间，用来捕捉KDA在不同范围内可能存在的非线性变化。",
    94: "阵容维度包含role_type、role_factor和role_match。role_type采用四类角色的one-hot编码，role_factor使用前文设定的角色权重，role_match标记当前角色是否适合该地图。该组特征用于把角色职责和地图适配关系输入模型。",
    95: "地图维度包含map_factor和map_match。map_factor使用第三章的地图权重，map_match根据localStorage中玩家历史战绩统计其高频地图，并标记当前地图是否属于前三张常用地图。该特征用于近似表示玩家的地图熟练度。",
    97: "训练集和测试集通过scikit-learn的train_test_split函数划分，测试集比例设为0.2，即300条。实验启用stratify参数按胜负标签分层抽样，保证训练集和测试集的胜负比例接近一致；随机种子固定为42，便于复现实验结果。",
    100: "模型对比在相同训练/测试划分下进行。GBDT、随机森林和逻辑回归使用同一组共享特征，为保证公平，对比时去除级联特征elo_pred，只保留其余11个输入变量。表4-1列出三个模型在准确率、精确率、召回率和F1分数上的结果。",
    103: "实验结果显示，GBDT在四项指标上均领先。其准确率为85.2%，高于逻辑回归的76.0%，也高于随机森林的81.0%。这一结果与模型特性相符：Boosting串行迭代能够持续修正前一轮难分样本，适合处理胜负概率接近的对局。级联特征elo_pred的贡献也在后续实验中单独分析。",
    108: "消融实验表明，三个修正因子的贡献并不相同。移除KDA修正后准确率下降2.1个百分点，说明临场表现是最直接的解释变量；移除角色修正下降1.3个百分点，反映角色职责差异具有独立作用；移除地图修正下降0.8个百分点，幅度较小，但仍说明地图适应性会影响预测结果。",
    114: "用户认证模块用于解决多用户数据隔离和功能访问控制。它由login.html登录页面、app.py中的Token鉴权接口和index.html中的前端路由守卫组成。用户登录后，前端携带Token访问后续功能，后端据此区分不同用户的数据范围。",
    115: "登录页采用电竞风格视觉设计。页面以深色背景和粒子动画营造氛围，中心卡片使用半透明背景、模糊滤镜和渐变边框，并在卡片内部加入扫描线动画。登录和注册表单放在同一卡片内，用户可通过切换按钮在两种模式之间转换。",
    120: "战绩导入模块提供原始数据输入。该模块主要由前端驱动，并把浏览器作为轻量数据持久化位置，减少高频写入对后端和SQLite的压力。录入表单包含比赛结果、地图、Kills/Deaths/Assists、角色定位和对局时间等字段，其中Deaths默认值为1，用于避免KDA计算时出现除零问题。",
    121: "战绩统计面板会在数据变化后重算总场次、总胜率和平均KDA。角色维度胜率柱状图按D/C/S/I分组展示表现差异，地图分布饼图展示八张地图的对局占比。用户可以据此观察自己常用角色、常玩地图和整体胜率变化。",
    126: "模块内置五组典型场景参数，用于快速展示算法差异。高KDA胜利、低KDA胜利、高KDA失败、均势对局和高分差对局分别对应不同评分变化场景。用户不需要逐项录入参数，也能观察标准ELO和优化ELO在相同条件下的更新差别。",
    131: "历史记录模块统一查询ELO计算、胜率预测和对比实验三类数据。GET /api/history/all接口一次性读取elo_history、predictions和comparisons三张表的最近记录，并组合成JSON返回给前端。用户可在同一页面回看不同功能产生的历史结果，也可删除单条记录。",
    133: "数据可视化仪表盘以ECharts渲染图表，并通过GET /api/dashboard/stats获取当前用户的数据汇总。页面包含ELO评分趋势、评分变化分布、胜负比例和地图/角色统计等组件。折线图从elo_history中读取评分时间序列，帮助用户观察个人竞技水平的长期变化。",
    137: "ELO天梯排行榜由服务端/api/leaderboard接口聚合生成。后端遍历users表，查询每个用户最新ELO记录，并统计其总场次和胜场数，再按评分排序返回。该接口把分数、胜率、场次和用户信息整合到排行榜数据中，供前端统一展示。",
    138: "排行榜前端采用分层展示。顶部用领奖台布局突出前三名，中部表格展示全服排名、玩家名、段位、ELO、胜率、总场次和KDA等字段，当前用户用高亮样式标识。该设计兼顾竞技仪式感和数据查询效率。",
    143: "每个大段位的ELO区间再划分为三个子级别。系统根据玩家ELO落入的区间确定大段位，再通过线性插值计算其在子级别中的位置和积分。这样既能保留ELO连续分值的精度，也能把结果转化为玩家更熟悉的段位展示。",
    147: "阵容推荐模块使用多因子规则引擎，根据localStorage中积累的历史战绩，为当前地图给出角色定位建议。用户至少需要3场历史战绩后才能启用推荐，保证计算具有基础样本。系统分别为D/C/S/I四类角色计算得分，再选出最高分角色。",
    148: "推荐评分包含地图角色适配度、该角色历史胜率、该角色平均KDA、近期状态和团队需求等维度。地图适配度来自pMapRoleBonus静态矩阵，例如Icebox中发起位因侦察技能优势获得较高分。近期状态使用指数衰减权重，使最近对局比早期对局影响更大。",
    149: "推荐结果分三层展示。顶部给出角色类型缩写和置信度，中部列出推荐理由，并附带该角色在当前地图的胜率、近期KDA等支撑数据；底部用提示卡说明推荐仅作为辅助参考。这样的结构让用户先看到结论，再查看依据。",
    165: "前端整体采用Vue.js 2.7 CDN运行时和单页HTML组织方式。系统没有引入Node.js构建链，Vue逻辑、CSS和DOM模板都写在index.html中。页面通过标签页导航和条件渲染切换七个功能区，适合小型演示系统快速部署和运行。",
    166: "除各功能标签页外，系统还提供全局段位选择弹窗和Toast提示组件。段位选择器由showTierSelector控制显示状态，Toast通过showToast(message,type)在不同模块中复用。ECharts图表在数据更新后重新渲染，避免切换页面后图表状态不同步。",
    168: "系统对外提供20个RESTful风格API，覆盖用户认证、ELO计算、胜率预测、算法对比、历史记录、排行榜、仪表盘统计和段位自定义等业务域。表6-1汇总了接口路径、HTTP方法、功能说明和认证需求，便于说明前后端交互边界。",
}


def iter_paragraphs(document: Document):
    """遍历正文段落和表格单元格段落，用于统一统计和空格清理。"""
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def paragraph_colors(paragraph) -> set[str]:
    """读取段落中所有显式字体颜色。"""
    colors: set[str] = set()
    for run in paragraph.runs:
        try:
            color = run.font.color.rgb if run.font.color and run.font.color.rgb else None
        except Exception:
            color = None
        if color:
            colors.add(str(color).upper())
    return colors


def replace_paragraph_text(paragraph, text: str) -> None:
    """替换热点段文字，保留段落样式，并用绿色标记已改写内容。"""
    first_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = GREEN
    if first_run is not None:
        run.bold = first_run.bold
        run.italic = first_run.italic
        run.underline = first_run.underline
        run.font.size = first_run.font.size
        run.font.name = first_run.font.name


def normalize_cjk_english_spacing(paragraph) -> int:
    """删除中文与英文字母之间的普通空格，保留目录制表符和标题编号空格。"""
    changes = 0
    text_nodes = paragraph._p.xpath(".//w:t")
    for node in text_nodes:
        old = node.text or ""
        new = LETTER_SPACE_RE.sub("", old)
        if new != old:
            node.text = new
            changes += 1

    for left, right in zip(text_nodes, text_nodes[1:]):
        left_text = left.text or ""
        right_text = right.text or ""
        new_left = re.sub(r"(?<=[\u4e00-\u9fff]) +$", "", left_text) if re.match(r"^[A-Za-z]", right_text) else left_text
        new_left = re.sub(r"(?<=[A-Za-z]) +$", "", new_left) if re.match(r"^[\u4e00-\u9fff]", right_text) else new_left
        new_right = re.sub(r"^ +(?=[A-Za-z])", "", right_text) if re.search(r"[\u4e00-\u9fff]$", new_left) else right_text
        new_right = re.sub(r"^ +(?=[\u4e00-\u9fff])", "", new_right) if re.search(r"[A-Za-z]$", new_left) else new_right
        if new_left != left_text:
            left.text = new_left
            changes += 1
        if new_right != right_text:
            right.text = new_right
            changes += 1
    return changes


def count_document(document: Document) -> dict[str, int]:
    """统计验证所需的字符数、热点颜色和绿色段落。"""
    texts: list[str] = []
    green_count = 0
    green_han = 0
    risk_count = 0
    for paragraph in iter_paragraphs(document):
        text = paragraph.text
        texts.append(text)
        colors = paragraph_colors(paragraph)
        if colors & RISK_COLORS:
            risk_count += 1
        if "1F8F3A" in colors:
            green_count += 1
            green_han += len(HAN_RE.findall(text))
    all_text = "\n".join(texts)
    return {
        "total_han_chars": len(HAN_RE.findall(all_text)),
        "risk_color_paragraphs": risk_count,
        "green_paragraphs": green_count,
        "green_han_chars": green_han,
        "cjk_english_letter_spaces": len(LETTER_SPACE_RE.findall(all_text)),
    }


def write_report(before: dict[str, int], after: dict[str, int], backup_path: Path, skipped: list[int]) -> None:
    """生成单独改稿说明，说明热点范围和验证结果。"""
    lines = [
        "# AIGC 改稿说明",
        "",
        f"- 工作稿：{DOCX_PATH.resolve()}",
        f"- 输出稿：{FINAL_DOCX.resolve()}",
        f"- 备份稿：{backup_path.resolve()}",
        f"- 本轮处理热点段落数：{after['green_paragraphs']}",
        f"- 跳过段落序号：{skipped if skipped else '无'}",
        "",
        "## 本轮处理的热点",
        "",
        "- 红色高风险段：处理图神经网络相关研究和数据可视化综述中的模板化表述。",
        "- 橙色中风险段：处理ELO原理、VALORANT机制、模型评估、特征工程、实验结果和系统实现描述。",
        "- 紫色提示段：处理KDA修正、战绩导入、历史记录、数据仪表盘和阵容推荐等实现段落。",
        "",
        "## 改写策略",
        "",
        "- 保留报告已有事实、引用编号、算法名称、接口路径和实验数值，不新增无法追溯的项目内容。",
        "- 将泛化综述句改为更直接的任务说明，将过长并列句拆成短句，减少机械化连接词。",
        "- 保留原图表、表格、公式和章节位置，没有重画旧图表。",
        "- 被改写段落统一标为绿色，便于和检测报告原色标注区分。",
        "",
        "## 验证结果",
        "",
        f"- 中文字符数（不含英文、标点）：{before['total_han_chars']} -> {after['total_han_chars']}",
        f"- 原红/橙/紫热点段落数：{before['risk_color_paragraphs']} -> {after['risk_color_paragraphs']}",
        f"- 绿色改写段落数：{after['green_paragraphs']}",
        f"- 绿色热点段中文字符数：{after['green_han_chars']}",
        f"- 中文与英文字母之间空格匹配数：{before['cjk_english_letter_spaces']} -> {after['cjk_english_letter_spaces']}",
        "",
        "## 下轮复检建议",
        "",
        "- 重新检测后只针对剩余红色或橙色页继续改，不建议整篇重写。",
    ]
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = WORK_DIR / f"working_valorant_aigc_report_backup_{timestamp}.docx"
    shutil.copy2(DOCX_PATH, backup_path)

    document = Document(str(DOCX_PATH))
    before = count_document(document)
    skipped: list[int] = []

    for index, replacement in REWRITES.items():
        if index >= len(document.paragraphs) or not document.paragraphs[index].text.strip():
            skipped.append(index)
            continue
        replace_paragraph_text(document.paragraphs[index], replacement)

    spacing_changes = 0
    for paragraph in iter_paragraphs(document):
        spacing_changes += normalize_cjk_english_spacing(paragraph)

    after = count_document(document)
    document.save(DOCX_PATH)
    shutil.copy2(DOCX_PATH, FINAL_DOCX)
    write_report(before, after, backup_path, skipped)

    print(f"backup={backup_path}")
    print(f"spacing_node_changes={spacing_changes}")
    print(f"before={before}")
    print(f"after={after}")
    print(f"final={FINAL_DOCX}")
    print(f"report={REPORT_PATH}")


if __name__ == "__main__":
    main()
